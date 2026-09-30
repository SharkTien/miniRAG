"""Build the Vietnamese milestone and data-research report as a PDF."""

from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.graphics.shapes import Drawing, Line, Polygon, Rect, String
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    HRFlowable,
    Image,
    KeepTogether,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.platypus.tableofcontents import TableOfContents


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "milestones_bonus_ec.pdf"


def register_fonts() -> tuple[str, str]:
    """Register a Unicode font available on the build host."""
    candidates = (
        (
            "/home/ntcai/rag-blueprint/parse-nemotron/NeMo-Retriever/cache/vietnamese-ocr/site/matplotlib/mpl-data/fonts/ttf/DejaVuSans.ttf",
            "/home/ntcai/rag-blueprint/parse-nemotron/NeMo-Retriever/cache/vietnamese-ocr/site/matplotlib/mpl-data/fonts/ttf/DejaVuSans-Bold.ttf",
        ),
        (
            "/home/ntcai/face-reconizer/face-recognition/.venv311/lib/python3.11/site-packages/matplotlib/mpl-data/fonts/ttf/DejaVuSans.ttf",
            "/home/ntcai/face-reconizer/face-recognition/.venv311/lib/python3.11/site-packages/matplotlib/mpl-data/fonts/ttf/DejaVuSans-Bold.ttf",
        ),
        (
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        ),
        (
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf",
        ),
    )
    for regular, bold in candidates:
        if Path(regular).exists() and Path(bold).exists():
            pdfmetrics.registerFont(TTFont("ReportSans", regular))
            pdfmetrics.registerFont(TTFont("ReportSans-Bold", bold))
            return "ReportSans", "ReportSans-Bold"
    raise FileNotFoundError("No Unicode TrueType font was found")


class ReportDocument(BaseDocTemplate):
    """Document template with page numbers and an automatic table of contents."""

    def __init__(self, filename: str, regular: str, **kwargs):
        super().__init__(filename, **kwargs)
        frame = Frame(
            self.leftMargin,
            self.bottomMargin,
            self.width,
            self.height,
            id="normal",
        )
        self.addPageTemplates(
            [PageTemplate(id="main", frames=frame, onPage=self._draw_page)]
        )
        self.regular = regular

    def _draw_page(self, canvas, document):
        canvas.saveState()
        canvas.setFont(self.regular, 8)
        canvas.setFillColor(colors.HexColor("#64748B"))
        if document.page > 1:
            canvas.drawString(18 * mm, 12 * mm, "Mini RAG Service | Hồ sơ bàn giao")
            canvas.drawRightString(
                A4[0] - 18 * mm, 12 * mm, f"Trang {document.page}"
            )
            canvas.setStrokeColor(colors.HexColor("#CBD5E1"))
            canvas.line(18 * mm, 16 * mm, A4[0] - 18 * mm, 16 * mm)
        canvas.restoreState()

    def afterFlowable(self, flowable):
        if not isinstance(flowable, Paragraph):
            return
        level = {"ReportH1": 0, "ReportH2": 1, "ReportH3": 2}.get(
            flowable.style.name
        )
        if level is not None:
            text = flowable.getPlainText()
            key = f"h{level}-{self.page}-{id(flowable)}"
            self.canv.bookmarkPage(key)
            self.canv.addOutlineEntry(text, key, level=level, closed=False)
            self.notify("TOCEntry", (level, text, self.page))


def p(text: str, style: ParagraphStyle) -> Paragraph:
    """Create a paragraph with safe XML text."""
    return Paragraph(escape(text).replace("\n", "<br/>").replace("  ", "&nbsp; "), style)


def rich(text: str, style: ParagraphStyle) -> Paragraph:
    """Create a paragraph for intentionally limited inline markup."""
    return Paragraph(text, style)


def evidence_image(path: Path, caption: str, regular: str, width: float = 166 * mm):
    """Return a bounded screenshot with a Vietnamese caption."""
    reader = ImageReader(str(path))
    image_width, image_height = reader.getSize()
    height = width * image_height / image_width
    image = Image(str(path), width=width, height=height)
    image.hAlign = "CENTER"
    caption_style = ParagraphStyle(
        "ImageCaption",
        fontName=regular,
        fontSize=8,
        leading=10.5,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#475569"),
        spaceBefore=3,
        spaceAfter=9,
    )
    return KeepTogether([image, Paragraph(escape(caption), caption_style)])


def pipeline_diagram(regular: str) -> Drawing:
    """Draw the detailed ingestion and retrieval architecture used in the report."""
    drawing = Drawing(510, 320)
    palette = {
        "data": colors.HexColor("#FDE3A7"),
        "process": colors.HexColor("#BFD7F7"),
        "control": colors.HexColor("#16324A"),
        "output": colors.HexColor("#A9E4C5"),
        "ink": colors.HexColor("#16324A"),
        "blue": colors.HexColor("#2563EB"),
        "line": colors.HexColor("#CBD5E1"),
    }

    def label(x, y, text, size=7.4, color="#16324A", bold=False, anchor="start"):
        drawing.add(String(x, y, text, fontName="ReportSans-Bold" if bold else regular,
                           fontSize=size, fillColor=colors.HexColor(color), textAnchor=anchor))

    def box(x, y, w, h, title, lines, fill, title_color="#16324A"):
        drawing.add(Rect(x, y, w, h, rx=9, ry=9, fillColor=fill,
                         strokeColor=colors.HexColor("#64748B"), strokeWidth=0.8))
        label(x + 10, y + h - 18, title, 9.2, title_color, True)
        for idx, line in enumerate(lines):
            label(x + 10, y + h - 33 - idx * 11, line, 6.8, "#334155")

    def arrow(x1, y1, x2, y2, dashed=False):
        line = Line(x1, y1, x2, y2, strokeColor=palette["blue"], strokeWidth=1.25)
        if dashed:
            line.strokeDashArray = [4, 3]
        drawing.add(line)
        import math
        angle = math.atan2(y2 - y1, x2 - x1)
        size = 5
        left = (x2 - size * math.cos(angle - 0.45), y2 - size * math.sin(angle - 0.45))
        right = (x2 - size * math.cos(angle + 0.45), y2 - size * math.sin(angle + 0.45))
        drawing.add(Polygon([x2, y2, left[0], left[1], right[0], right[1]],
                            fillColor=palette["blue"], strokeColor=None))

    # Header and legend, matching the reference atlas style.
    label(4, 292, "MINI RAG / PIPELINE ATLAS", 8.5, "#2563EB", True)
    label(4, 274, "Kiến trúc luồng dữ liệu và truy xuất", 17, "#16324A", True)
    legend = [("Dữ liệu", palette["data"]), ("Xử lý", palette["process"]),
              ("Điều phối / kiểm soát", palette["control"]), ("Đầu ra", palette["output"])]
    lx = 5
    for name, fill in legend:
        drawing.add(Rect(lx, 251, 11, 11, rx=2, ry=2, fillColor=fill,
                         strokeColor=colors.HexColor("#64748B"), strokeWidth=0.5))
        label(lx + 16, 253, name, 7.1)
        lx += 100 if name != "Điều phối / kiểm soát" else 137

    # Ingestion lane: source data becomes searchable evidence.
    label(5, 227, "LUỒNG NẠP DỮ LIỆU", 7.2, "#176B87", True)
    top = [(5, "Tài liệu", ["PDF · DOCX · TXT"], palette["data"]),
           (105, "Trích xuất", ["Docling · OCR NVIDIA", "giữ trang / bố cục"], palette["process"]),
           (205, "Chuẩn hóa", ["làm sạch · chia đoạn", "bảng / hình / chú thích"], palette["process"]),
           (305, "Vector hóa", ["embedding NVIDIA", "metadata + nguồn"], palette["process"]),
           (405, "Kho bằng chứng", ["PostgreSQL + pgvector", "file gốc ở kho đối tượng"], palette["data"])]
    for x, title, lines, fill in top:
        box(x, 170, 90, 47, title, lines, fill)
    for x in [95, 195, 295, 395]:
        arrow(x, 193, x + 10, 193)

    # Query lane and control plane.
    label(5, 145, "LUỒNG TRUY VẤN", 7.2, "#176B87", True)
    bottom = [(5, "Câu hỏi", ["ý định · thực thể", "ràng buộc cần trả lời"], palette["data"]),
              (105, "Lập kế hoạch", ["từ khóa quan trọng", "phạm vi tài liệu"], palette["process"]),
              (205, "Hybrid search", ["BM25 + dense", "mở rộng ứng viên"], palette["process"]),
              (305, "Xếp hạng", ["lọc tài liệu", "khả năng trả lời"], palette["process"]),
              (405, "Answer + source", ["mô hình NVIDIA", "nguồn · chunk · trang"], palette["output"])]
    for x, title, lines, fill in bottom:
        box(x, 84, 90, 47, title, lines, fill)
    for x in [95, 195, 295, 395]:
        arrow(x, 107, x + 10, 107)
    # Shared storage and the control plane.
    arrow(450, 170, 450, 132)
    label(455, 151, "đọc bằng chứng", 6.2, "#475569")
    drawing.add(Rect(178, 140, 154, 24, rx=6, ry=6, fillColor=palette["control"], strokeColor=palette["control"]))
    label(255, 154, "ĐIỀU PHỐI RAG", 8.2, "#FFFFFF", True, "middle")
    label(255, 145, "phạm vi · điểm · nguồn", 6.3, "#D8E7F5", False, "middle")
    arrow(150, 131, 178, 151)
    arrow(332, 151, 355, 131)
    arrow(350, 84, 332, 151, dashed=True)

    # Explanatory panel.
    drawing.add(Rect(5, 10, 500, 60, rx=7, ry=7, fillColor=colors.white,
                     strokeColor=palette["line"], strokeWidth=0.8))
    label(14, 54, "ĐIỂM CHÍNH", 7.4, "#008C95", True)
    label(14, 41, "Tài liệu trở thành bằng chứng có nguồn trước khi mô hình trả lời.", 7.8)
    label(14, 28, "Ví dụ: câu hỏi về NAPAS phải giữ được đoạn 2–5 ngày và mốc tính tiền.", 7.4, "#475569")
    label(326, 54, "CÁCH ĐỌC", 7.4, "#2563EB", True)
    label(326, 41, "Nét liền = luồng chính · nét đứt = phản hồi", 7.0, "#475569")
    label(326, 28, "Kho bằng chứng dùng chung cho truy xuất và trích nguồn.", 7.0, "#475569")
    drawing.add(Line(5, 78, 505, 78, strokeColor=palette["line"], strokeWidth=0.6))
    label(5, 4, "Luồng nạp dữ liệu ở trên · luồng hỏi đáp ở dưới · lớp điều phối kiểm soát nguồn", 6.8, "#64748B")
    return drawing


def architecture_diagram(regular: str) -> Drawing:
    """Draw the deployable architecture and model orchestration layers."""
    drawing = Drawing(510, 292)
    navy = colors.HexColor("#16324A")
    blue = colors.HexColor("#2563EB")
    pale_blue = colors.HexColor("#BFD7F7")
    yellow = colors.HexColor("#FDE3A7")
    green = colors.HexColor("#A9E4C5")
    line_color = colors.HexColor("#CBD5E1")

    def text(x, y, value, size=7.2, color="#16324A", bold=False, anchor="start"):
        drawing.add(String(x, y, value, fontName="ReportSans-Bold" if bold else regular,
                           fontSize=size, fillColor=colors.HexColor(color), textAnchor=anchor))

    def block(x, y, w, h, title, detail, fill, title_color="#16324A"):
        drawing.add(Rect(x, y, w, h, rx=8, ry=8, fillColor=fill,
                         strokeColor=colors.HexColor("#64748B"), strokeWidth=0.8))
        text(x + 9, y + h - 17, title, 8.4, title_color, True)
        detail_color = "#D8E7F5" if title_color == "#FFFFFF" else "#334155"
        for n, line in enumerate(detail):
            text(x + 9, y + h - 31 - n * 10, line, 6.6, detail_color)

    def connect(x1, y1, x2, y2, dashed=False):
        ln = Line(x1, y1, x2, y2, strokeColor=blue, strokeWidth=1.15)
        if dashed:
            ln.strokeDashArray = [4, 3]
        drawing.add(ln)
        import math
        angle = math.atan2(y2 - y1, x2 - x1)
        left = (x2 - 5 * math.cos(angle - 0.45), y2 - 5 * math.sin(angle - 0.45))
        right = (x2 - 5 * math.cos(angle + 0.45), y2 - 5 * math.sin(angle + 0.45))
        drawing.add(Polygon([x2, y2, left[0], left[1], right[0], right[1]],
                            fillColor=blue, strokeColor=None))

    text(4, 277, "MINI RAG / ARCHITECTURE ATLAS", 8.5, "#2563EB", True)
    text(4, 259, "Các lớp triển khai và mô hình xử lý", 17, "#16324A", True)
    # Layer labels at the left.
    text(4, 231, "GIAO DIỆN", 7, "#176B87", True)
    text(4, 181, "API", 7, "#176B87", True)
    text(4, 126, "LƯU TRỮ", 7, "#176B87", True)
    text(4, 70, "MÔ HÌNH", 7, "#176B87", True)
    # Main horizontal architecture.
    block(62, 207, 105, 36, "Frontend / CLI", ["upload · chat · streaming"], yellow)
    block(188, 207, 105, 36, "FastAPI", ["documents · query · health"], pale_blue)
    block(314, 207, 105, 36, "Nginx", ["giao diện production"], pale_blue)
    connect(167, 225, 188, 225)
    connect(293, 225, 314, 225)
    block(62, 151, 105, 42, "Document API", ["validate · hash", "queued document"], pale_blue)
    block(188, 151, 105, 42, "Query API", ["plan · retrieve", "stream answer"], pale_blue)
    block(314, 151, 105, 42, "Worker queue", ["claim · retry", "cancel on delete"], navy, "#FFFFFF")
    connect(114, 207, 114, 193)
    connect(240, 207, 240, 193)
    connect(366, 207, 366, 193)
    block(62, 95, 105, 42, "MinIO", ["file gốc · crop", "snapshot trang"], yellow)
    block(188, 95, 105, 42, "PostgreSQL", ["state · metadata", "conversation"], yellow)
    block(314, 95, 105, 42, "pgvector", ["embedding · index", "full-text"], yellow)
    connect(114, 151, 114, 137)
    connect(240, 151, 240, 137)
    connect(366, 151, 366, 137)
    connect(293, 116, 314, 116)
    # Model pipeline below.
    block(62, 39, 105, 37, "Extract + OCR", ["Docling · NVIDIA OCR"], pale_blue)
    block(188, 39, 105, 37, "Embedding", ["NVIDIA Embed 1B"], pale_blue)
    block(314, 39, 105, 37, "Grounded LLM", ["NVIDIA NIM · sources"], green)
    connect(114, 95, 114, 76)
    connect(240, 95, 240, 76)
    connect(366, 95, 366, 76)
    connect(167, 57, 188, 57)
    connect(293, 57, 314, 57)
    # Retrieval control feedback.
    connect(240, 151, 240, 137, dashed=True)
    text(247, 143, "lọc / xếp hạng", 6.2, "#475569")
    drawing.add(Rect(5, 5, 500, 23, rx=6, ry=6, fillColor=colors.white,
                     strokeColor=line_color, strokeWidth=0.8))
    text(14, 19, "ĐIỂM CHÍNH", 7.2, "#008C95", True)
    text(83, 19, "API nhận yêu cầu; worker xử lý nặng; kho lưu trạng thái và bằng chứng; mô hình chỉ nhận ngữ cảnh đã lọc.", 6.8, "#475569")
    return drawing

def table(data, widths, regular, header=True, small=False):
    font_size = 7.3 if small else 8.1
    leading = 9.2 if small else 10.2
    converted = []
    for row_index, row in enumerate(data):
        converted.append(
            [
                Paragraph(
                    escape(str(value)),
                    ParagraphStyle(
                        "Cell",
                        fontName="ReportSans-Bold" if header and row_index == 0 else regular,
                        fontSize=font_size,
                        leading=leading,
                        textColor=colors.HexColor("#0F172A"),
                    ),
                )
                for value in row
            ]
        )
    result = Table(converted, colWidths=widths, repeatRows=1 if header else 0)
    commands = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
    ]
    if header:
        commands.extend(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F3D5E")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ]
        )
    for row_index in range(1 if header else 0, len(data)):
        if row_index % 2:
            commands.append(
                ("BACKGROUND", (0, row_index), (-1, row_index), colors.HexColor("#F8FAFC"))
            )
    result.setStyle(TableStyle(commands))
    return result


def build():
    regular, _ = register_fonts()
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="CoverTitle",
            parent=styles["Title"],
            fontName="ReportSans-Bold",
            fontSize=25,
            leading=31,
            textColor=colors.HexColor("#0F3D5E"),
            alignment=TA_CENTER,
            spaceAfter=12,
        )
    )
    styles.add(
        ParagraphStyle(
            name="CoverSub",
            parent=styles["Normal"],
            fontName=regular,
            fontSize=12,
            leading=17,
            textColor=colors.HexColor("#475569"),
            alignment=TA_CENTER,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ReportH1",
            parent=styles["Heading1"],
            fontName="ReportSans-Bold",
            fontSize=16,
            leading=20,
            textColor=colors.HexColor("#0F3D5E"),
            spaceBefore=13,
            spaceAfter=8,
            keepWithNext=True,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ContentsTitle",
            parent=styles["ReportH1"],
            fontName="ReportSans-Bold",
        )
    )
    styles.add(
        ParagraphStyle(
            name="ReportH2",
            parent=styles["Heading2"],
            fontName="ReportSans-Bold",
            fontSize=11.5,
            leading=15,
            textColor=colors.HexColor("#176B87"),
            spaceBefore=9,
            spaceAfter=5,
            keepWithNext=True,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ReportH3",
            parent=styles["Heading3"],
            fontName="ReportSans-Bold",
            fontSize=9.5,
            leading=12,
            textColor=colors.HexColor("#334155"),
            spaceBefore=7,
            spaceAfter=3,
            keepWithNext=True,
        )
    )
    styles.add(
        ParagraphStyle(
            name="BodyVi",
            parent=styles["BodyText"],
            fontName=regular,
            fontSize=9.2,
            leading=13.2,
            textColor=colors.HexColor("#1E293B"),
            spaceAfter=6,
        )
    )
    styles.add(
        ParagraphStyle(
            name="BulletVi",
            parent=styles["BodyText"],
            fontName=regular,
            fontSize=8.9,
            leading=12.2,
            leftIndent=13,
            firstLineIndent=-8,
            bulletIndent=0,
            textColor=colors.HexColor("#1E293B"),
            spaceAfter=3,
        )
    )
    styles.add(
        ParagraphStyle(
            name="TOCVi",
            parent=styles["Normal"],
            fontName=regular,
            fontSize=10,
            leading=14,
            leftIndent=14,
            firstLineIndent=-14,
            textColor=colors.HexColor("#1E293B"),
        )
    )
    styles.add(
        ParagraphStyle(
            name="Callout",
            parent=styles["BodyText"],
            fontName=regular,
            fontSize=9.2,
            leading=13,
            borderColor=colors.HexColor("#93C5FD"),
            borderWidth=0.6,
            borderPadding=8,
            backColor=colors.HexColor("#EFF6FF"),
            textColor=colors.HexColor("#0F172A"),
            spaceBefore=5,
            spaceAfter=8,
        )
    )
    styles.add(
        ParagraphStyle(
            name="EvalText",
            parent=styles["BodyText"],
            fontName=regular,
            fontSize=8.1,
            leading=10.4,
            textColor=colors.HexColor("#1E293B"),
            spaceAfter=0,
        )
    )

    document = ReportDocument(
        str(OUTPUT),
        regular,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=17 * mm,
        bottomMargin=23 * mm,
        title="Báo cáo milestone và nghiên cứu dữ liệu Mini RAG Service",
        author="Mini RAG Service",
    )
    toc = TableOfContents()
    toc.levelStyles = [
        ParagraphStyle("TOC0", parent=styles["TOCVi"], leftIndent=0, fontName="ReportSans-Bold"),
        ParagraphStyle("TOC1", parent=styles["TOCVi"], leftIndent=12, fontSize=9),
        ParagraphStyle("TOC2", parent=styles["TOCVi"], leftIndent=24, fontSize=8.5),
    ]

    story = [
        Spacer(1, 32 * mm),
        rich("MINI RAG / HỒ SƠ KỸ THUẬT", styles["CoverSub"]),
        Spacer(1, 7 * mm),
        rich("Luồng dữ liệu và truy xuất tài liệu", styles["CoverTitle"]),
        rich("Các mốc triển khai, bằng chứng kiểm thử và kết quả đánh giá", styles["CoverSub"]),
        Spacer(1, 15 * mm),
        HRFlowable(width="65%", thickness=2, color=colors.HexColor("#F59E0B"), hAlign="CENTER"),
        Spacer(1, 15 * mm),
        rich("Ngày cập nhật: 30/09/2026", styles["CoverSub"]),
        rich('<link href="https://github.com/SharkTien/miniRAG" color="#176B87">Kho mã nguồn trên GitHub: github.com/SharkTien/miniRAG</link>', styles["CoverSub"]),
        Spacer(1, 30 * mm),
        PageBreak(),
        rich("Mục lục", styles["ContentsTitle"]),
        toc,
        PageBreak(),
    ]

    def h1(text):
        story.append(Paragraph(escape(text), styles["ReportH1"]))

    def h2(text):
        story.append(Paragraph(escape(text), styles["ReportH2"]))

    def body(text):
        story.append(p(text, styles["BodyVi"]))

    def bullet(text):
        story.append(Paragraph("• " + escape(text), styles["BulletVi"]))

    h1("3. Các mốc thực hiện")
    h2("Mốc 1 — Luồng dữ liệu")
    body("Đã hoàn thành tải tài liệu, kiểm tra định dạng, lưu kho đối tượng, trích xuất PDF/TXT/DOCX và định dạng mở rộng, nhận dạng ký tự NVIDIA cho bản quét và ảnh, làm sạch, chuẩn hóa ngữ nghĩa có giới hạn thời gian, chia đoạn, tạo vector theo lô, lưu pgvector và xử lý nền.")
    story.append(pipeline_diagram(regular))
    body("Sơ đồ trên là luồng xử lý dữ liệu cuối cùng: tài liệu được trích xuất trước, làm sạch và chia đoạn theo cấu trúc, sau đó mới tạo vector và lưu cùng metadata. Khi hỏi, hệ thống dùng đồng thời điểm từ khóa và điểm ngữ nghĩa, lọc theo tài liệu, xếp hạng bằng chứng rồi mới gửi ngữ cảnh phù hợp cho mô hình trả lời.")
    h3 = lambda text: story.append(Paragraph(escape(text), styles["ReportH3"]))
    h3("Các mô hình và vai trò trong luồng")
    story.append(table([
        ["Thành phần", "Mô hình/công nghệ", "Vai trò"],
        ["Nhận dạng ký tự", "NVIDIA NeMo Retriever OCR v2", "Đọc chữ trong PDF quét, ảnh và vùng bảng khi tài liệu không có lớp chữ."],
        ["Tạo vector", "NVIDIA NeMo Retriever Embed 1B", "Biến câu hỏi và từng đoạn thành vector cùng không gian để tìm kiếm ngữ nghĩa."],
        ["Chuẩn hóa ngữ nghĩa", "Mô hình ngôn ngữ qua NVIDIA NIM", "Chuẩn hóa tiêu đề, quan hệ bảng và tín hiệu truy vấn khi cần; có giới hạn thời gian và thử lại."],
        ["Sinh câu trả lời", "Mô hình ngôn ngữ qua NVIDIA NIM", "Tổng hợp câu trả lời chỉ từ các bằng chứng đã được chọn và trả kèm nguồn."],
        ["Lưu trữ và tìm kiếm", "PostgreSQL + pgvector + toàn văn", "Lưu vector/metadata và kết hợp tìm kiếm ngữ nghĩa với từ khóa."],
    ], [38 * mm, 57 * mm, 83 * mm], regular, small=True))
    story.append(PageBreak())
    h3("Sơ đồ kiến trúc triển khai")
    story.append(architecture_diagram(regular))
    body("Kiến trúc tách rõ bốn lớp: giao diện gửi yêu cầu, API điều phối, lưu trữ/tiến trình nền và lớp mô hình. Tải tài liệu chỉ ghi nhận nhanh trạng thái queued; worker mới thực hiện trích xuất, OCR, chuẩn hóa, chia đoạn và tạo vector. Truy vấn đi qua planner, hybrid retrieval và EvidenceService trước khi mô hình ngôn ngữ nhận context.")
    bullet("Hồ sơ nghiên cứu dữ liệu được bàn giao kèm báo cáo và dùng làm căn cứ cho các quyết định kỹ thuật.")
    bullet("877 câu hỏi được phân loại theo bảng, biểu mẫu, hình, chú thích và văn bản.")
    bullet("Các phương án OCR, mở rộng tập ứng viên, tìm kiếm kết hợp, xử lý bảng, hình và chú thích được chạy trên cùng bộ câu hỏi.")
    h2("Mốc 2 — API hỏi đáp")
    body("API hỏi đáp, lập kế hoạch truy vấn, tìm kiếm toàn văn kết hợp vector, mở rộng tập ứng viên, lọc danh tính tài liệu, xếp hạng bằng chứng, lời nhắc bám nguồn, trích dẫn nguồn và trả kết quả từng phần đã có trong mã nguồn.")
    h2("Mốc 3 — Đóng gói và kiểm tra tự động")
    body("Dockerfile, Docker Compose cho API, tiến trình nền, PostgreSQL + pgvector, MinIO và giao diện, tệp cấu hình mẫu và quy trình GitHub Actions đã có. Cổng chất lượng đạt 98,8/100; ảnh chụp quy trình nằm ở phần bằng chứng.")
    h2("Mốc 4 — Tài liệu và trình diễn")
    body("README, tài liệu kiến trúc, ví dụ API, kịch bản trình diễn, truy vết yêu cầu, bảng kiểm tra truy xuất và báo cáo PDF này được bàn giao.")

    h1("4. Bằng chứng trực quan từ hệ thống")
    body("Các ảnh dưới đây được lấy từ quá trình chạy thật của hệ thống. Chúng minh họa đường đi của dữ liệu và cấu trúc lưu trữ, không chỉ là sơ đồ khái niệm.")
    story.append(evidence_image(ROOT / "output/images/docker_CICD.png", "Hình 1. Quy trình kiểm tra tự động: các bước chất lượng, giao diện và dựng Docker hoàn tất với trạng thái thành công trong cùng một lần chạy.", regular))
    story.append(evidence_image(ROOT / "output/images/pgvector.png", "Hình 2. Cấu trúc PostgreSQL/pgvector: bảng document_chunks có content dạng text, metadata dạng jsonb, embedding vector(2048), khóa liên kết document_id và chỉ mục theo tài liệu.", regular))
    story.append(evidence_image(ROOT / "output/images/upload.png", "Hình 3. Giao diện kho tài liệu sau khi tải lên: sáu tài liệu đã xuất hiện, tài liệu Shopee được chọn, bản xem trước PDF và phần văn bản đã định dạng được hiển thị cùng số lượng chunk.", regular, width=166 * mm))
    story.append(evidence_image(ROOT / "output/images/chunk.png", "Hình 4. Màn hình xem chunk: nội dung được chia thành các đoạn có thứ tự, giữ tiêu đề, trang và văn bản liên quan để bước truy xuất có thể chọn đúng phần.", regular, width=130 * mm))

    h1("5. Nghiên cứu và xử lý dữ liệu")
    h2("5.1 Hồ sơ dữ liệu")
    body("Benchmark gồm 877 câu hỏi trên 5 PDF. Bảng chiếm 299 câu, biểu mẫu 107, hình 159, chú thích/đánh dấu 285 và văn bản thường 27. Việc phân loại giúp không dùng một chiến lược chunk/OCR cho mọi loại nội dung.")
    story.append(table([
        ["Loại dữ liệu", "Số câu", "Rủi ro cần kiểm soát"],
        ["Bảng", "299", "Mất quan hệ hàng, cột và tiêu đề"],
        ["Biểu mẫu", "107", "Nhãn và giá trị bị tách rời"],
        ["Hình", "159", "Dữ kiện nằm trong sơ đồ/hình"],
        ["Chú thích", "285", "Callout, tô màu, vùng đánh dấu"],
        ["Văn bản", "27", "Header/footer và chunk quá ngắn"],
    ], [35 * mm, 20 * mm, 105 * mm], regular, small=True))
    h2("5.2 Thiết kế phép đo")
    bullet("Document hit: tài liệu đúng xuất hiện trong kết quả.")
    bullet("Content hit: kết quả có đoạn chứa dữ kiện trả lời.")
    bullet("Assertion pass: câu trả lời đúng mệnh đề được kiểm tra.")
    body("Ba chỉ số được đọc cùng nhau. Document hit cao nhưng content hit thấp nghĩa là hệ thống biết tên tài liệu nhưng chưa tìm được bằng chứng; content hit có nhưng assertion pass thấp chỉ ra vấn đề ở bước tổng hợp hoặc diễn đạt.")
    h2("5.3 So sánh phương án")
    story.append(table([
        ["Phương án", "Kết quả", "Quyết định"],
        ["Ban đầu", "85,52% document; 33,52% content; 18,59% pass", "Mốc"],
        ["OCR toàn trang", "34,44% content; 19,73% pass", "Giữ"],
        ["OCR + ứng viên x8", "88,26% document; 37,66% content", "Giữ"],
        ["Hybrid dense + lexical", "90,65% document; 41,33% content; 23,72% pass", "Baseline"],
        ["Structured table toàn cục", "Content còn 36,85%", "Loại"],
        ["Visual element + region OCR", "Figure 34/159 hit; 19/159 pass", "Giữ cho hình"],
        ["OCR tile cố định", "Annotation 92/285 hit, thấp hơn tuyến ngữ cảnh", "Loại"],
    ], [53 * mm, 72 * mm, 35 * mm], regular, small=True))
    h2("5.4 Quyết định kỹ thuật từ dữ liệu")
    bullet("Tài liệu quét: OCR toàn trang NVIDIA có timeout và fallback để không treo job.")
    bullet("Bảng: giữ quan hệ tiêu đề–hàng–cột, thêm chunk cha và điều hướng theo câu hỏi.")
    bullet("Hình: lưu crop, page, bbox và metadata; benchmark cho thấy chỉ mục phần tử riêng tốt hơn gắn ảnh vào chunk trang.")
    bullet("Chú thích: không mặc định dùng OCR lát cắt cố định; cần detector theo hình học và quan hệ vùng đánh dấu.")
    bullet("Truy xuất: kết hợp dense và lexical, mở rộng tập ứng viên rồi mới rerank theo khả năng trả lời.")

    h1("6. Kiểm thử đơn vị và kiểm thử API")
    body("Lệnh kiểm tra chính là ruff check app tests, python -m compileall -q app tests và python -m pytest -q tests. Kết quả gần nhất: 22 kiểm thử đạt; kiểm thử gọi dịch vụ NVIDIA trực tiếp được bỏ qua khi môi trường không có khóa API, còn các kiểm thử ngoại tuyến vẫn chạy bình thường.")
    h2("6.1 Kiểm thử thành phần lõi")
    body("Mục tiêu: bảo đảm dữ liệu đầu vào được làm sạch và chia đúng trước khi tạo vector. Mẫu kiểm tra: chuỗi có khoảng trắng thừa, ba dòng xuống dòng liên tiếp và một ký tự điều khiển phải trở thành một dòng trống đúng chuẩn; đoạn văn có số trang và tiêu đề phải giữ lại trang/phần; hộp giới hạn phải được chuẩn hóa về cùng hệ tọa độ. Nếu một ký tự điều khiển làm thay đổi nội dung hoặc mất số trang, kiểm thử thất bại.")
    h2("6.2 Kiểm thử hợp đồng API")
    body("Mục tiêu: xác nhận dịch vụ phản hồi đúng khi người dùng gọi API. Mẫu kiểm tra trạng thái hoạt động: gửi GET /health và yêu cầu mã phản hồi thành công cùng trạng thái dịch vụ cơ sở dữ liệu. Mẫu hỏi đáp: gửi POST /query với câu hỏi hợp lệ và kiểm tra có câu trả lời cùng nguồn. Mẫu lỗi: câu hỏi rỗng phải trả lỗi 400; tệp ngoài PDF/TXT/DOCX phải bị từ chối; tệp hợp lệ phải trả trạng thái queued và mã tài liệu; truy vấn không có bằng chứng phải trả sources rỗng, không dựng nguồn giả.")
    h2("6.3 Kiểm thử xếp hạng bằng chứng")
    body("Mục tiêu: đoạn chứa dữ kiện quyết định phải đứng trên đoạn chỉ trùng chủ đề. Mẫu kiểm tra: với câu hỏi về thời gian hoàn tiền thẻ NAPAS, đoạn có “2–5 ngày làm việc” phải vượt đoạn chỉ nói thời hạn trả hàng; đổi tên tệp nhưng giữ nguyên nội dung không được làm thay đổi điểm nội dung. Kiểm thử cũng xác nhận đoạn Kamereo, Microsoft hoặc chính sách bảo hành khác không lọt vào khi câu hỏi chỉ hỏi Shopee.")
    h2("6.4 Kiểm thử nhận dạng ký tự và định tuyến")
    body("Mục tiêu: chọn đúng cách đọc theo loại tài liệu. Mẫu kiểm tra: PDF đã có lớp chữ phải đi đường tắt, PDF quét phải chuyển sang OCR NVIDIA, dấu tiếng Việt phải được sửa khi OCR trả ký tự lỗi, vùng bảng phải giữ hộp giới hạn và hệ thống phải dừng đúng thời gian khi mô hình OCR chậm. Văn bản sau OCR phải được đánh dấu phương pháp trích xuất và độ tin cậy để có thể kiểm tra lại.")
    h2("6.5 Kiểm thử kết nối mô hình NVIDIA")
    body("Mục tiêu: kiểm tra cấu hình, chia lô và khả năng thử lại mà không bắt buộc gọi dịch vụ trực tuyến trong mọi lần chạy. Mẫu kiểm tra: đọc đúng mô hình, chia trang thành các lô có kích thước giới hạn, giữ thứ tự trang sau khi xử lý song song, và từ chối câu trả lời không có bằng chứng. Khi thiếu khóa truy cập, phần gọi trực tuyến được bỏ qua có chủ đích; các kiểm thử ngoại tuyến vẫn phải đạt.")
    h2("6.6 Kiểm chứng truy xuất")
    body("Mục tiêu: kiểm tra độc lập việc tạo vector và tìm đúng đoạn. Mẫu kiểm tra gồm các câu hỏi về thời hạn hoàn tiền, phí trả hàng, chính sách bảo hành, hotline và xử lý hàng lỗi; mỗi câu lưu lại tài liệu mong đợi, tài liệu nhận được, câu trả lời mong đợi và câu trả lời thực tế để đánh giá ở mục 8.")
    body("Kiểm thử API dùng dịch vụ giả lập để kiểm tra hợp đồng phản hồi: câu trả lời phải đi kèm sources; câu hỏi rỗng phải bị từ chối; truy vấn không có bằng chứng không được sinh nguồn giả; file không hỗ trợ phải trả lỗi; upload hợp lệ phải trả trạng thái queued.")

    h1("7. Cấu hình và nguồn sự thật")
    body("Cấu hình được tách khỏi mã nghiệp vụ. Tệp cấu hình mẫu chỉ chứa giá trị mẫu; giá trị thật được nạp từ môi trường và không đưa vào mã nguồn. Bộ đọc cấu hình kiểm tra kiểu dữ liệu và cung cấp giá trị mặc định cho máy phát triển.")
    story.append(table([
        ["Nhóm", "Biến cấu hình", "Tác dụng"],
        ["Kết nối", "DATABASE_URL, POSTGRES_DB, POSTGRES_USER", "Kết nối PostgreSQL/pgvector và tên cơ sở dữ liệu"],
        ["Mô hình", "EMBEDDING_MODEL, EMBEDDING_DIM, LLM_MODEL", "Mô hình tạo vector, số chiều và mô hình sinh"],
        ["Chia đoạn", "CHUNK_SIZE, CHUNK_OVERLAP", "Kích thước đoạn và phần chồng lấn"],
        ["Truy xuất", "TOP_K, RETRIEVAL_CANDIDATE_MULTIPLIER, SIMILARITY_THRESHOLD", "Số bằng chứng, tập ứng viên và ngưỡng lọc"],
        ["OCR", "NVIDIA_OCR_MODEL, NVIDIA_OCR_TIMEOUT_SECONDS, NVIDIA_OCR_BATCH_SIZE", "Mô hình, thời gian chờ và kích thước lô OCR"],
        ["Chuẩn hóa", "NIM_MODEL, NIM_REQUEST_TIMEOUT_SECONDS, NIM_MAX_RETRIES", "Mô hình ngữ nghĩa, thời gian chờ và số lần thử lại"],
    ], [35 * mm, 78 * mm, 63 * mm], regular, small=True))
    body("Khi triển khai, chỉ cần thay các biến môi trường và khởi động lại dịch vụ. Luồng xử lý không cần sửa mã nguồn để đổi mô hình, kích thước đoạn, số lượng kết quả hay thời gian chờ.")

    h1("8. Đánh giá từng trường hợp kiểm tra truy xuất")
    body("Bảng kiểm tra truy xuất có 8 câu hỏi có dữ liệu; hai dòng cuối để trống nên không được tính là trường hợp kiểm tra. Cách đánh giá: Đạt khi câu trả lời chứa đủ ý chính và không mâu thuẫn; Đạt một phần khi có ý chính nhưng thiếu điều kiện hoặc thêm thông tin chưa có căn cứ; Chưa đạt khi mâu thuẫn với đáp án hoặc bỏ sót dữ kiện quyết định.")
    # Đọc nguyên văn bộ kiểm tra để không làm mất câu hỏi hoặc chi tiết câu trả lời.
    workbook = openpyxl.load_workbook(ROOT / "output/retrieval_verification.xlsx", data_only=True)
    sheet = workbook.active
    source_labels = {
        "shopee": "Shopee — hướng dẫn đổi trả và hoàn tiền",
        "adore": "Adore — chính sách đổi hàng",
        "gd100": "DAT Group — chính sách bảo hành sản phẩm năng lượng mặt trời",
        "inverter": "DAT Group — chính sách bảo hành sản phẩm năng lượng mặt trời",
        "kamereo": "Kamereo — chính sách đổi trả hàng hóa",
    }
    verdicts = ["Đạt", "Đạt", "Chưa đạt", "Chưa đạt", "Đạt", "Chưa đạt", "Đạt", "Đạt"]

    def readable_source(source):
        folded = str(source or "").lower()
        for key, label in source_labels.items():
            if key in folded:
                return label
        return "Nguồn trong bộ kiểm tra"

    def exact(value):
        return escape(str(value or "")).replace("\n", "<br/>")

    retrieval_cases = []
    for row in list(sheet.iter_rows(min_row=2, values_only=True)):
        question, expected_source, retrieved_source, expected_answer, retrieved_answer = row[:5]
        if not question:
            continue
        retrieval_cases.append((
            str(question),
            readable_source(expected_source),
            readable_source(retrieved_source),
            str(expected_answer or ""),
            str(retrieved_answer or ""),
        ))

    comparison_notes = [
        ("Khớp nội dung cốt lõi:", "2–5 ngày làm việc và hoàn về thẻ NAPAS đều xuất hiện đúng.", "#15803D"),
        ("Khớp nội dung cốt lõi:", "Ba phương thức trả hàng và điều kiện hỗ trợ cước được nêu đầy đủ.", "#15803D"),
        ("Sai lệch chính:", "Hệ thống khẳng định được tách riêng áo và hoàn tiền, trong khi ground truth giữ kết luận mở.", "#B91C1C"),
        ("Sai lệch chính:", "Hệ thống khẳng định được đổi phụ kiện, trong khi ground truth yêu cầu nêu thiếu dữ kiện.", "#B91C1C"),
        ("Khớp nội dung cốt lõi:", "Thời hạn 1,5 năm và mốc ngày giao hàng đều đúng; phần thêm không làm đổi đáp án.", "#15803D"),
        ("Sai lệch chính:", "Hệ thống bỏ mốc 24 giờ và quy trình báo giá khi máy nứt vỡ.", "#B91C1C"),
        ("Khớp nội dung cốt lõi:", "Số hotline 0812 46 37 27 trùng ground truth.", "#15803D"),
        ("Khớp nội dung cốt lõi:", "Không trừ công nợ ngay và đổi/bù ở đơn tiếp theo đều được trả đúng.", "#15803D"),
    ]
    for index, (question, expected_source, retrieved_source, expected, observed) in enumerate(retrieval_cases):
        verdict = verdicts[index] if index < len(verdicts) else "Chưa phân loại"
        source_match = "Khớp" if expected_source == retrieved_source else "Không khớp"
        note_title, note_text, note_color = comparison_notes[index]
        verdict_color = "#15803D" if verdict == "Đạt" else "#B91C1C"
        header = Table([[Paragraph(f"<b>CA {index + 1:02d}</b>  ·  {escape(verdict)}", ParagraphStyle(
            f"EvalHeader{index}", parent=styles["BodyText"], fontName="ReportSans-Bold", fontSize=9.2,
            leading=11, textColor=colors.white)), Paragraph(escape("Đối chiếu ground truth ↔ output"), ParagraphStyle(
            f"EvalHeaderRight{index}", parent=styles["BodyText"], fontName=regular, fontSize=7.8,
            leading=10, alignment=TA_RIGHT, textColor=colors.white))]], colWidths=[34 * mm, 136 * mm])
        header.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(verdict_color)),
            ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        card = [Spacer(1, 5), header]
        card.append(rich(f"<b><font color=\"#2563EB\">CÂU HỎI</font></b><br/>{exact(question)}", styles["EvalText"]))
        card.append(Table([            [rich('<b><font color="#15803D">GROUND TRUTH</font></b>', styles["EvalText"]), rich('<b><font color="#2563EB">OUTPUT HỆ THỐNG</font></b>', styles["EvalText"])],
            [rich(f"<b>Nguồn mong đợi:</b> {exact(expected_source)}<br/><br/><b>Đáp án mong đợi:</b><br/>{exact(expected)}", styles["EvalText"]),
             rich(f"<b>Nguồn hệ thống:</b> {exact(retrieved_source)}<br/><b>Đối chiếu nguồn:</b> {source_match}<br/><br/><b>Câu trả lời hệ thống:</b><br/>{exact(observed)}", styles["EvalText"])],
        ], colWidths=[85 * mm, 85 * mm], style=TableStyle([
            ("BACKGROUND", (0, 0), (0, 0), colors.HexColor("#DCFCE7")),
            ("BACKGROUND", (1, 0), (1, 0), colors.HexColor("#DBEAFE")),
            ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#F8FAFC")),
            ("BOX", (0, 0), (-1, -1), 0.45, colors.HexColor("#CBD5E1")),
            ("INNERGRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ])))
        card.append(rich(f"<b><font color=\"{note_color}\">SO SÁNH:</font></b> <font color=\"{note_color}\"><b>{escape(note_title)}</b></font> {escape(note_text)}", styles["EvalText"]))
        story.append(KeepTogether(card))
    
    body("Kết quả tổng hợp: 5/8 câu đạt đầy đủ và 3/8 câu chưa đạt. Cả 8 trường hợp đều truy xuất đúng tài liệu mong đợi; ba lỗi còn lại nằm ở mức độ đầy đủ và tính kỷ luật của câu trả lời, không phải chọn sai tài liệu.")

    h1("9. Cách tạo source trong câu trả lời")
    body("Source không do mô hình sinh tự đặt ra. Dịch vụ hỏi đáp lấy source từ những chunk đã vượt qua truy xuất kết hợp, lọc danh tính tài liệu, phân bố điểm và xếp hạng bằng chứng. Với mỗi chunk được chọn, hệ thống tạo một khối bằng chứng có tên file, mã chunk và trang trước khi gọi mô hình.")
    story.append(table([
        ["Trường trong source", "Ý nghĩa"],
        ["file_name, document_id", "Tài liệu và bản ghi tài liệu gốc"],
        ["chunk_id, page, page_start, page_end", "Đoạn và vị trí có thể mở lại để kiểm tra"],
        ["dense_score, bm25_score, similarity_score", "Điểm vector, điểm từ khóa và điểm tương đồng"],
        ["query_compatibility_score, answerability_score", "Mức phù hợp với câu hỏi và khả năng chứa dữ kiện trả lời"],
        ["source_locator, element_ids, section", "Vị trí, phần tử bố cục và mục nội dung"],
        ["extraction_method, ocr_confidence, snippet", "Cách lấy dữ liệu, độ tin cậy OCR và đoạn xem nhanh"],
    ], [62 * mm, 116 * mm], regular, small=True))
    body("Trong prompt, bằng chứng được đánh dấu dạng [EVIDENCE #... | File... | Chunk... | Trang...]. Sau khi mô hình sinh câu trả lời, API trả lại answer cùng mảng sources; nếu không có bằng chứng thì sources rỗng và hệ thống không dựng nguồn giả.")

    h1("10. Hạn chế và kế hoạch tiếp theo")
    bullet("Bộ đánh giá hiện tập trung 5 PDF SynthDocQA; cần thêm bộ tài liệu thực tế cân bằng theo loại dữ liệu.")
    bullet("Chỉ mục phần tử hình và OCR vùng chú thích chưa bật cho mọi tài liệu vận hành.")
    bullet("Bộ xếp hạng hiện dựa trên luật; chưa có mô hình xếp hạng học từ phản hồi người dùng.")
    bullet("Cần nạp lại tài liệu, chạy đánh giá hồi quy và theo dõi độ trễ, tỷ lệ OCR dự phòng, tỷ lệ không có bằng chứng theo loại dữ liệu.")
    body("Chi tiết phương pháp, phân tích lỗi, bảng kết quả đầy đủ và quy trình tái lập nằm trong hồ sơ nghiên cứu dữ liệu. File PDF này là bản trình bày có dẫn chứng; hồ sơ nghiên cứu là nơi lưu giải thích đầy đủ.")

    document.multiBuild(story)
    print(f"Wrote {OUTPUT} ({OUTPUT.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
